from datetime import date

from openpyxl import load_workbook

from src import cfg, sheet
from src.models import Account, Score


def test_export_xlsx_roundtrip():
    out = cfg.ROOT / "cache" / "_xlsx_test" / "w.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    a = Account(id="a", name="Acme", city="Pune", segment="corporate", first_seen=date(2026, 3, 1))
    sc = Score(account_id="a", fit=20, trigger=25, reach=0, total=45, reason="new campus", rank=1)
    sheet.export_xlsx(out, {"Accounts": [a], "Scores": [sc], "Drafts": []})
    wb = load_workbook(out)
    assert wb.sheetnames == ["Accounts", "Scores", "Drafts"]
    ws = wb["Accounts"]
    assert [c.value for c in ws[1]][:3] == ["id", "name", "domain"] and ws["B2"].value == "Acme" and ws.freeze_panes == "A2"
    assert wb["Scores"]["E2"].value == 45 and wb["Drafts"].max_row == 1       # empty tab still has its header
    out.unlink()
