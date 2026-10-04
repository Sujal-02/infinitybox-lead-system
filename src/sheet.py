"""Storage back-ends: Google Sheets through a service-account key or the Apps Script gateway, or a local .xlsx workbook.
Both sheet back-ends expose the same few worksheet calls, so the pipeline does not care which one is used."""
import json

import gspread
from pydantic import BaseModel

from .cfg import env
from .models import Account, Draft, Lead, LeadRow, Person, Score, Signal, Source

TABS: dict[str, type[BaseModel]] = {
    "Accounts": Account, "Signals": Signal, "People": Person,
    "Scores": Score, "Drafts": Draft, "Leads": Lead, "Sources": Source, "Pipeline": LeadRow,
}


def columns(tab: str) -> list[str]:
    return list(TABS[tab].model_fields)


class ScriptWS:
    """The few worksheet calls the pipeline uses, sent to the Apps Script gateway (inbound/apps_script.gs, sheetOp)."""
    def __init__(self, book, title):
        self.book, self.title = book, title

    def row_values(self, n):
        v = self.book.call("values", self.title) or []
        return v[n - 1] if len(v) >= n else []

    def append_row(self, r):
        self.append_rows([r])

    def append_rows(self, rows):
        for i in range(0, len(rows), 200):
            self.book.call("append", self.title, rows=rows[i:i + 200])

    def get_all_records(self):
        v = self.book.call("values", self.title)
        return [dict(zip(v[0], r)) for r in v[1:]] if v else []

    def clear(self):
        self.book.call("clear", self.title)


class ScriptBook:
    def __init__(self, url, token):
        self.url, self.token = url, token

    def call(self, op, title="", **kw):
        import requests
        r = requests.post(self.url, data=json.dumps({"op": op, "token": self.token, "title": title, **kw}), headers={"Content-Type": "text/plain"}, timeout=60)
        j = r.json()
        if not j.get("ok"):
            raise RuntimeError(f"sheet gateway: {j.get('error', r.status_code)}")
        return j.get("values") if op == "values" else j

    def worksheets(self):
        return [ScriptWS(self, t) for t in self.call("list")["titles"]]

    def worksheet(self, title):
        return ScriptWS(self, title)

    def add_worksheet(self, title, rows=0, cols=0):
        self.call("add", title)

    @property
    def title(self):
        return self.call("list")["title"]


def enabled() -> bool:
    """A Google Sheet is used when there is a service-account key or an Apps Script gateway; otherwise a local workbook is written."""
    from .cfg import ROOT
    return bool(env("SHEET_WEBHOOK_URL")) or (ROOT / env("GOOGLE_CREDS_PATH", "creds.json")).exists()


def open_book():
    if env("SHEET_WEBHOOK_URL"):
        return ScriptBook(env("SHEET_WEBHOOK_URL"), env("SHEET_API_TOKEN"))
    gc = gspread.service_account(filename=env("GOOGLE_CREDS_PATH", "creds.json"))
    return gc.open_by_key(env("GSHEET_ID"))


EVENTS_HEADER = ["ts", "sid", "event", "utm_source", "utm_medium", "utm_campaign", "ref", "path"]  # written by inbound/apps_script.gs


def check() -> str:
    """Open the sheet and say what is wrong in plain words if it cannot be opened. Returns a short success line."""
    from .cfg import ROOT
    if env("SHEET_WEBHOOK_URL"):
        try:
            book = open_book()
            return f"ok: '{book.title}' opened through the Apps Script gateway; tabs: {', '.join(w.title for w in book.worksheets())}"
        except Exception as e:
            raise SystemExit(f"gateway failed: {e}. Check SHEET_WEBHOOK_URL (the /exec URL), SHEET_API_TOKEN equals the script's API_TOKEN, and the web app is deployed with access: Anyone")
    path = ROOT / env("GOOGLE_CREDS_PATH", "creds.json")
    if not path.exists():
        raise SystemExit(f"{path.name} not found: download the service-account JSON key from Google Cloud and save it here")
    if not env("GSHEET_ID"):
        raise SystemExit("GSHEET_ID is empty: put the id from the sheet's URL (between /d/ and /edit) in .env")
    who = json.loads(path.read_text("utf-8")).get("client_email", "?")
    try:
        book = open_book()
    except gspread.exceptions.SpreadsheetNotFound:
        raise SystemExit(f"cannot open the sheet. Share it (Editor) with {who} and check GSHEET_ID")
    except gspread.exceptions.APIError as e:
        raise SystemExit(f"Google refused: {str(e)[:200]}. Enable the Google Sheets API and the Google Drive API for the project of {who}")
    return f"ok: '{book.title}' opened as {who}; tabs: {', '.join(w.title for w in book.worksheets())}"


def init(book) -> list[str]:
    """Create any missing tabs with header rows. Returns the tabs created."""
    have = {w.title for w in book.worksheets()}
    made = []
    for tab in TABS:
        if tab not in have:
            book.add_worksheet(title=tab, rows=1000, cols=len(columns(tab)))
            made.append(tab)
        ws = book.worksheet(tab)
        if not ws.row_values(1):
            ws.append_row(columns(tab))
    if "Events" not in have:  # anonymous page visits, filled by the Apps Script
        book.add_worksheet(title="Events", rows=1000, cols=len(EVENTS_HEADER))
        book.worksheet("Events").append_row(EVENTS_HEADER)
        made.append("Events")
    return made


def _row(tab: str, m: BaseModel) -> list:
    d = m.model_dump(mode="json")
    return ["" if d[c] is None else d[c] for c in columns(tab)]


def append(book, tab: str, rows: list[BaseModel]) -> None:
    if rows:
        book.worksheet(tab).append_rows([_row(tab, r) for r in rows])


def read(book, tab: str) -> list[BaseModel]:
    recs = book.worksheet(tab).get_all_records()
    return [TABS[tab].model_validate({k: v for k, v in r.items() if v != ""}) for r in recs]


def replace(book, tab: str, rows: list[BaseModel]) -> None:
    """Overwrite a tab's data (used for Scores, which is recomputed each run)."""
    ws = book.worksheet(tab)
    ws.clear()
    ws.append_rows([columns(tab)] + [_row(tab, r) for r in rows])


def export_xlsx(path, tables: dict[str, list[BaseModel]]) -> None:
    """Local alternative to Google Sheets: one tab per table, header row, filters, frozen header. No credentials needed.
    Open in Excel, or drag into Google Drive / import into a Google Sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    wb.remove(wb.active)
    for tab, rows in tables.items():
        ws = wb.create_sheet(tab)
        cols = columns(tab)
        ws.append(cols)
        for r in rows:
            ws.append(_row(tab, r))
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cols, 1):
            width = min(max([len(str(col))] + [len(str(v)) for v in (ws.cell(row=j, column=i).value for j in range(2, min(ws.max_row, 60) + 1))]) + 2, 60)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
            for j in range(2, ws.max_row + 1):
                ws.cell(row=j, column=i).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
